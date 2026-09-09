"""DocumentAssembler: reconstructs the translated DOCX from translated chunks.

- Orders chunks strictly by chunk_index (never alphabetically).
- Validates completeness BEFORE generating: missing / failed / empty / duplicated
  chunks abort assembly with a clear error.
- Produces right-to-left Arabic output with proper fonts, alignment and
  complex-script run properties.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.core.logging_config import get_logger
from app.models.chunk import TranslationChunk
from app.services.document.structure import (
    Block,
    HEADING,
    LIST_ITEM,
    PARAGRAPH,
    TABLE,
    blocks_from_json,
)

logger = get_logger(__name__)

ARABIC_FONT = "Arial"
ARABIC_CS_FONT = "Traditional Arabic"
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
HEADING_SIZES = {1: 20, 2: 16, 3: 14, 4: 12}
BODY_SIZE = 11

_MARKER_RE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)")


class AssemblyError(Exception):
    """Raised when the document cannot be safely assembled."""


def _set_paragraph_rtl(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.append(bidi)


def _style_run_rtl(run, size: int = BODY_SIZE) -> None:
    rPr = run._r.get_or_add_rPr()
    if rPr.find(qn("w:rtl")) is None:
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ARABIC_FONT)
    rFonts.set(qn("w:hAnsi"), ARABIC_FONT)
    rFonts.set(qn("w:cs"), ARABIC_CS_FONT)
    szCs = rPr.find(qn("w:szCs"))
    if szCs is None:
        szCs = OxmlElement("w:szCs")
        rPr.append(szCs)
    szCs.set(qn("w:val"), str(size * 2))


def _add_rich_runs(paragraph, text: str, size: int = BODY_SIZE, bold=False, italic=False) -> None:
    """Add runs honoring **bold**, *italic*, ***bold-italic*** markers."""
    for part in _MARKER_RE.split(text):
        if not part:
            continue
        run_bold, run_italic = bold, italic
        content = part
        if part.startswith("***") and part.endswith("***") and len(part) > 6:
            content = part[3:-3]
            run_bold = run_italic = True
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            content = part[2:-2]
            run_bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            content = part[1:-1]
            run_italic = True
        run = paragraph.add_run(content)
        run.bold = run_bold
        run.italic = run_italic
        run.font.size = Pt(size)
        _style_run_rtl(run, size)


def _set_table_rtl(table) -> None:
    tblPr = table._tbl.tblPr
    if tblPr.find(qn("w:bidiVisual")) is None:
        bidi = OxmlElement("w:bidiVisual")
        tblPr.append(bidi)


def _split_translated_paragraphs(translation: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", translation.strip()) if p.strip()]


def _map_translated(blocks: list[Block], translated: list[str]) -> list[str | None]:
    """Map translated paragraphs onto source blocks without ever losing text.

    - exact count -> 1:1
    - tables -> pipe-containing paragraphs are structurally matched to table
      blocks first (their content is ' | ' rows)
    - remaining prose -> in-order when counts match, otherwise distributed
      proportionally: each paragraph goes to the source block whose cumulative
      character range contains the paragraph's midpoint.
    """
    n = len(blocks)
    if n == 0:
        return []
    if not translated:
        return [None] * n
    if len(translated) == n:
        return list(translated)

    result: list[str | None] = [None] * n

    # --- structural pass: pipe paragraphs belong to table blocks ---
    table_indexes = [i for i, b in enumerate(blocks) if b.type == TABLE and b.rows]
    pipe_paras = [t for t in translated if "|" in t]
    prose_paras = [t for t in translated if "|" not in t]
    for i, tbl_i in enumerate(table_indexes):
        if i < len(pipe_paras):
            result[tbl_i] = pipe_paras[i]
    # leftover pipe paragraphs (no matching table) rejoin the prose pool
    prose_paras = prose_paras + pipe_paras[len(table_indexes):]

    content_indexes = [i for i in range(n) if result[i] is None]

    # --- prose pass ---
    if len(prose_paras) <= len(content_indexes):
        for i, ci in enumerate(content_indexes):
            result[ci] = prose_paras[i] if i < len(prose_paras) else None
        return result

    # More prose paragraphs than content blocks: proportional distribution.
    src_lens = [max(1, len(blocks[ci].plain_text())) for ci in content_indexes]
    src_total = sum(src_lens)
    tr_lens = [max(1, len(t)) for t in prose_paras]
    tr_total = sum(tr_lens)

    bounds: list[tuple[float, float]] = []
    acc = 0.0
    for length in src_lens:
        start = acc / src_total
        acc += length
        bounds.append((start, acc / src_total))

    assignments: list[list[str]] = [[] for _ in content_indexes]
    tr_acc = 0.0
    for para, length in zip(prose_paras, tr_lens):
        mid = (tr_acc + length / 2) / tr_total
        tr_acc += length
        for i, (start, end) in enumerate(bounds):
            if start <= mid < end:
                assignments[i].append(para)
                break
        else:
            assignments[-1].append(para)

    # Every content block should receive at least one paragraph when possible.
    for i in range(len(content_indexes)):
        if not assignments[i]:
            for j in range(i + 1, len(content_indexes)):
                if assignments[j]:
                    assignments[i].append(assignments[j].pop(0))
                    break

    for i, ci in enumerate(content_indexes):
        group = assignments[i]
        result[ci] = "\n\n".join(group) if group else None
    return result


def _parse_table_lines(text: str, rows: list[list[str]]) -> list[list[str]]:
    """Parse translated pipe-rows back into a cell matrix of the same shape."""
    # Safety: strip any [n] row markers left by the translation prompt.
    text = re.sub(r"^\[[0-9\u0660-\u0669]{1,4}\]\s*", "", text, flags=re.MULTILINE)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    result: list[list[str]] = []
    for i, line in enumerate(lines):
        cells = [c.strip() for c in line.split("|")]
        # Strip empty edge cells created by leading/trailing pipes.
        if cells and cells[0] == "":
            cells = cells[1:]
        if cells and cells[-1] == "":
            cells = cells[:-1]
        if i < len(rows):
            # Keep source column count: merge/split as needed.
            n_cols = len(rows[i])
            if len(cells) < n_cols:
                cells = cells + [""] * (n_cols - len(cells))
            elif len(cells) > n_cols:
                cells = cells[: n_cols - 1] + [" | ".join(cells[n_cols - 1 :])]
        result.append(cells)
    return result if result else rows


class DocumentAssembler:
    def validate_chunks(self, chunks: list[TranslationChunk]) -> None:
        """Fail fast on any completeness problem. Never produce a broken file."""
        if not chunks:
            raise AssemblyError("No chunks to assemble.")
        indexes = [c.chunk_index for c in chunks]
        expected = list(range(len(chunks)))
        if sorted(indexes) != expected:
            missing = sorted(set(expected) - set(indexes))
            raise AssemblyError(
                f"Chunk sequence is invalid. Missing chunk indexes: {missing[:10]}."
            )
        failed = [c for c in chunks if c.status == "failed"]
        if failed:
            idxs = [c.chunk_index for c in failed]
            raise AssemblyError(
                f"{len(failed)} chunk(s) failed and must be retried before generating "
                f"the document: chunk indexes {idxs[:10]}."
            )
        pending = [c for c in chunks if c.status != "completed"]
        if pending:
            idxs = [c.chunk_index for c in pending]
            raise AssemblyError(
                f"Some chunks are incomplete ({len(pending)}). The final Word document "
                f"cannot yet be generated. Pending chunk indexes: {idxs[:10]}."
            )
        empty = [c for c in chunks if not (c.translation or "").strip()]
        if empty:
            idxs = [c.chunk_index for c in empty]
            raise AssemblyError(
                f"{len(empty)} chunk(s) have empty translations: {idxs[:10]}."
            )
        if len(set(indexes)) != len(indexes):
            raise AssemblyError("Duplicate chunk indexes detected.")

    # ------------------------------------------------------------------
    def assemble(
        self,
        document_title: str,
        chunks: list[TranslationChunk],
        output_path: Path,
    ) -> Path:
        self.validate_chunks(chunks)
        ordered = sorted(chunks, key=lambda c: c.chunk_index)

        doc = DocxDocument()

        # Document title (RTL).
        title_para = doc.add_paragraph()
        _set_paragraph_rtl(title_para)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(document_title or "Translated Document")
        title_run.bold = True
        title_run.font.size = Pt(22)
        title_run.font.color.rgb = HEADING_COLOR
        _style_run_rtl(title_run, 22)

        for chunk in ordered:
            blocks = blocks_from_json(chunk.source_blocks)
            if not blocks:
                blocks = [Block(seq=0, type=PARAGRAPH, text=chunk.source_text)]
            translated = _split_translated_paragraphs(chunk.translation or "")
            mapping = _map_translated(blocks, translated)
            for block, text in zip(blocks, mapping):
                if block.type == TABLE and block.rows:
                    # Keep only pipe rows (the translated table), skipping any
                    # prose paragraphs that landed in the same share.
                    pipe_lines = [l for l in (text or "").splitlines() if "|" in l]
                    table_text = "\n".join(pipe_lines)
                    rows = _parse_table_lines(table_text, block.rows) if table_text.strip() else block.rows
                    if len(rows) < 2:
                        rows = block.rows  # fallback: keep source shape
                    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                    table.style = "Table Grid"
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    _set_table_rtl(table)
                    for r, row in enumerate(rows):
                        for c in range(len(table.rows[r].cells)):
                            cell_text = row[c] if c < len(row) else ""
                            cell = table.rows[r].cells[c]
                            cell_para = cell.paragraphs[0]
                            _set_paragraph_rtl(cell_para)
                            _add_rich_runs(cell_para, cell_text, size=10, bold=(r == 0))
                    doc.add_paragraph()
                    continue

                if text is None or not text.strip():
                    # No translated text mapped to this block: the model merged
                    # its content into a neighboring paragraph. Skip rather
                    # than emitting English source text into the Arabic doc.
                    continue

                if block.type == HEADING:
                    level = min(4, max(1, block.level or 1))
                    para = doc.add_heading(level=level)
                    _set_paragraph_rtl(para)
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in list(para.runs):
                        run.text = ""
                    _add_rich_runs(para, text, size=HEADING_SIZES[level], bold=True)
                    for run in para.runs:
                        run.font.color.rgb = HEADING_COLOR
                elif block.type == LIST_ITEM:
                    para = doc.add_paragraph(style="List Bullet")
                    _set_paragraph_rtl(para)
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    _add_rich_runs(para, text, bold=block.bold, italic=block.italic)
                else:
                    para = doc.add_paragraph()
                    _set_paragraph_rtl(para)
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    _add_rich_runs(para, text, bold=block.bold, italic=block.italic)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(
            "DOCX assembled",
            extra={
                "operation": "assemble",
                "status": "success",
                "chunks": len(ordered),
                "output": output_path.name,
            },
        )
        return output_path
