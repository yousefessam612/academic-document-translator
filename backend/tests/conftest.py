"""Test configuration.

IMPORTANT: sets test env vars BEFORE any app import so the app binds to a
temporary database and storage directory.
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

TEST_ROOT = Path(tempfile.mkdtemp(prefix="adt-tests-"))
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_ROOT / 'test.db'}"
os.environ["STORAGE_DIR"] = str(TEST_ROOT / "storage")
os.environ["MAX_CONCURRENT_TRANSLATIONS"] = "2"
os.environ["AGENTROUTER_API_KEY"] = ""  # tests must never use real credentials
os.environ["AGENTROUTER_MODEL"] = ""

import pytest  # noqa: E402

from app.core.config import settings  # noqa: E402


@pytest.fixture(autouse=True)
def _clean_db():
    """Isolate every test: ensure directories and fresh tables."""
    settings.ensure_directories()
    from app.db.database import Base, SessionLocal, init_db

    init_db()
    session = SessionLocal()
    try:
        for table in reversed(Base.metadata.sorted_tables):
            session.execute(table.delete())
        session.commit()
    finally:
        session.close()
    yield


@pytest.fixture()
def db_session():
    from app.db.database import SessionLocal, init_db

    init_db()
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


# --------------------------------------------------------------- helpers
def make_blocks(n_paragraphs: int = 10, words: int = 60, with_headings: bool = True):
    from app.services.document.structure import Block

    blocks = []
    seq = 0
    for i in range(n_paragraphs):
        if with_headings and i % 5 == 0:
            level = 1 if i % 10 == 0 else 2
            blocks.append(Block(seq=seq, type="heading", text=f"Chapter {i//10 + 1}" if level == 1 else f"Section {i}", level=level, page=i + 1))
            seq += 1
        text = " ".join(f"word{j}" for j in range(words))
        blocks.append(Block(seq=seq, type="paragraph", text=text, page=i + 1))
        seq += 1
    return blocks


def make_text_pdf(path: Path, pages: int = 3) -> Path:
    import fitz

    doc = fitz.open()
    for p in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Chapter {p + 1} Visual Impairment", fontsize=22)
        y = 110
        for para in range(3):
            page.insert_text(
                (72, y),
                "Students with visual impairment need assistive technology. "
                * 3,
                fontsize=11,
            )
            y += 40
    doc.save(str(path))
    doc.close()
    return path


def make_scanned_pdf(path: Path, pages: int = 2) -> Path:
    """Image-only PDF (no text layer)."""
    import fitz

    # First render a text page to a pixmap.
    src = fitz.open()
    page = src.new_page()
    page.insert_text((72, 72), "Chapter 1 Low vision research", fontsize=18)
    page.insert_text((72, 110), "Orientation and mobility training matters.", fontsize=11)
    pix = page.get_pixmap(dpi=150)
    src.close()

    doc = fitz.open()
    for _ in range(pages):
        p = doc.new_page()
        p.insert_image(fitz.Rect(0, 0, 612, 792), pixmap=pix)
    doc.save(str(path))
    doc.close()
    return path


def make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Braille Literacy Study", level=1)
    doc.add_heading("Introduction", level=2)
    p = doc.add_paragraph()
    p.add_run("Visual impairment ").bold = True
    p.add_run("affects ")
    p.add_run("learning outcomes").italic = True
    doc.add_paragraph("Assistive technology includes screen readers and braille displays.")
    doc.add_paragraph("Key points:", style="List Bullet")
    doc.add_paragraph("Screen reader", style="List Bullet")
    doc.add_paragraph("Screen magnifier", style="List Bullet")
    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "Arabic"
    table.rows[0].cells[2].text = "Domain"
    table.rows[1].cells[0].text = "Low vision"
    table.rows[1].cells[1].text = "ضعف البصر"
    table.rows[1].cells[2].text = "Visual Impairment"
    doc.save(str(path))
    return path
