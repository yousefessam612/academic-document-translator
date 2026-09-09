"""DocumentAssembler tests: ordering, completeness validation, RTL DOCX."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from app.services.export.assembler import AssemblyError, DocumentAssembler


class FakeChunk:
    def __init__(self, index, source, translation, blocks, status="completed"):
        self.chunk_index = index
        self.source_text = source
        self.translation = translation
        self.source_blocks = blocks
        self.status = status


def heading(text, level=1):
    return {"seq": 0, "type": "heading", "text": text, "page": 1, "level": level,
            "marker": "", "rows": [], "bold": False, "italic": False}


def paragraph(text):
    return {"seq": 1, "type": "paragraph", "text": text, "page": 1, "level": 0,
            "marker": "", "rows": [], "bold": False, "italic": False}


class TestValidation:
    def test_missing_chunk_rejected(self, tmp_path):
        chunks = [
            FakeChunk(0, "a", "أ", [paragraph("a")]),
            FakeChunk(2, "c", "ج", [paragraph("c")]),  # chunk 1 missing
        ]
        with pytest.raises(AssemblyError, match="Missing chunk indexes"):
            DocumentAssembler().assemble("Book", chunks, tmp_path / "out.docx")

    def test_failed_chunk_rejected(self, tmp_path):
        chunks = [
            FakeChunk(0, "a", "أ", [paragraph("a")]),
            FakeChunk(1, "b", None, [paragraph("b")], status="failed"),
        ]
        with pytest.raises(AssemblyError, match="failed"):
            DocumentAssembler().assemble("Book", chunks, tmp_path / "out.docx")

    def test_empty_translation_rejected(self, tmp_path):
        chunks = [
            FakeChunk(0, "a", "أ", [paragraph("a")]),
            FakeChunk(1, "b", "  ", [paragraph("b")]),
        ]
        with pytest.raises(AssemblyError, match="empty"):
            DocumentAssembler().assemble("Book", chunks, tmp_path / "out.docx")

    def test_pending_chunk_rejected(self, tmp_path):
        chunks = [
            FakeChunk(0, "a", "أ", [paragraph("a")]),
            FakeChunk(1, "b", None, [paragraph("b")], status="pending"),
        ]
        with pytest.raises(AssemblyError):
            DocumentAssembler().assemble("Book", chunks, tmp_path / "out.docx")


class TestAssembly:
    def test_assemble_produces_rtl_docx(self, tmp_path: Path):
        blocks0 = [
            heading("Chapter One", 1),
            paragraph("First paragraph with **bold** and *italic*."),
        ]
        blocks1 = [
            {"seq": 5, "type": "table", "text": "", "page": 2, "level": 0, "marker": "",
             "rows": [["Term", "Arabic"], ["Low vision", "ضعف البصر"]], "bold": False, "italic": False},
            paragraph("Second chunk paragraph."),
        ]
        chunks = [
            FakeChunk(0, "Chapter One\n\nFirst paragraph.", "الفصل الأول\n\nالفقرة الأولى **بالخط العريض** و*المائل*.", blocks0),
            FakeChunk(1, "table\n\nSecond chunk paragraph.",
                      "المصطلح | الترجمة\nضعف البصر | Low vision\n\nالفقرة الثانية.", blocks1),
        ]
        out = tmp_path / "translated.docx"
        DocumentAssembler().assemble("Test Book", chunks, out)
        assert out.exists()

        doc = DocxDocument(str(out))
        # Title paragraph
        assert doc.paragraphs[0].text == "Test Book"

        # Heading preserved with level
        headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert any(h.text == "الفصل الأول" for h in headings)

        # RTL (bidi) set on paragraphs
        all_paras = doc.paragraphs + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        rtl_count = sum(
            1 for p in all_paras if p._p.find(qn("w:pPr")) is not None
            and p._p.find(qn("w:pPr")).find(qn("w:bidi")) is not None
        )
        assert rtl_count >= 5, "Arabic paragraphs must be right-to-left"

        # Table rebuilt with translated content
        assert len(doc.tables) == 1
        assert doc.tables[0].rows[0].cells[0].text == "المصطلح"
        assert "ضعف البصر" in doc.tables[0].rows[1].cells[0].text

        # Bold marker turned into actual bold run
        bold_runs = [r for p in doc.paragraphs for r in p.runs if r.bold]
        assert any("بالخط العريض" in r.text for r in bold_runs)

    def test_chunk_order_not_alphabetical(self, tmp_path):
        texts = [(i, f"para {i}") for i in range(5)]
        chunks = [
            FakeChunk(i, src, f"فقرة رقم {i}", [paragraph(src)])
            for i, src in reversed(texts)  # deliberately reversed input
        ]
        out = tmp_path / "order.docx"
        DocumentAssembler().assemble("Order Test", chunks, out)
        doc = DocxDocument(str(out))
        body_texts = [p.text for p in doc.paragraphs if p.text.startswith("فقرة")]
        assert body_texts == [f"فقرة رقم {i}" for i in range(5)]

    def test_source_blocks_fallback_plain(self, tmp_path):
        chunk = FakeChunk(0, "Just text", "نص فقط", None)
        out = tmp_path / "fallback.docx"
        DocumentAssembler().assemble("Fallback", [chunk], out)
        doc = DocxDocument(str(out))
        assert any(p.text == "نص فقط" for p in doc.paragraphs)


class TestParagraphDistribution:
    """The model often returns MORE paragraphs than source blocks (splitting
    merged run-on text). The assembler must never drop translated content."""

    def test_more_translated_than_blocks_no_content_loss(self, tmp_path):
        blocks = [
            heading("Chapter One", 1),
            paragraph("A" * 2000),  # long source paragraph
        ]
        # Model split the long paragraph into 5 translated paragraphs
        translation = (
            "الفصل الأول\n\n"
            "الفقرة الأولى كاملة هنا.\n\n"
            "الفقرة الثانية كاملة هنا.\n\n"
            "الفقرة الثالثة كاملة هنا.\n\n"
            "الفقرة الرابعة كاملة هنا.\n\n"
            "الفقرة الخامسة كاملة هنا."
        )
        chunk = FakeChunk(0, "Chapter One\n\nAAAA...", translation, blocks)
        out = tmp_path / "dist.docx"
        DocumentAssembler().assemble("Dist Test", [chunk], out)
        doc = DocxDocument(str(out))
        full = "\n".join(p.text for p in doc.paragraphs)
        # Heading mapped correctly
        assert any(p.text == "الفصل الأول" for p in doc.paragraphs if p.style.name.startswith("Heading"))
        # EVERY translated paragraph must survive in the output
        for i, word in enumerate(["الأولى", "الثانية", "الثالثة", "الرابعة", "الخامسة"]):
            assert f"الفقرة {word} كاملة هنا." in full, f"paragraph {i} was dropped"

    def test_fewer_translated_than_blocks_no_english_leak(self, tmp_path):
        blocks = [
            heading("Chapter Two", 1),
            paragraph("First source paragraph."),
            paragraph("Second source paragraph."),
        ]
        # Model merged everything into ONE paragraph (validation allows -1)
        translation = "الفصل الثاني\n\nالفقرة الأولى والثانية مدمجتان في فقرة عربية واحدة."
        chunk = FakeChunk(0, "Chapter Two\n\nFirst...\n\nSecond...", translation, blocks)
        out = tmp_path / "fewer.docx"
        DocumentAssembler().assemble("Fewer Test", [chunk], out)
        doc = DocxDocument(str(out))
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "مدمجتان" in full
        # The unmapped block must NOT leak English source text
        assert "Second source paragraph." not in full
        assert "First source paragraph." not in full

    def test_exact_match_unchanged(self, tmp_path):
        blocks = [heading("Chapter", 1), paragraph("Content here.")]
        translation = "الفصل\n\nالمحتوى هنا."
        chunk = FakeChunk(0, "Chapter\n\nContent here.", translation, blocks)
        out = tmp_path / "exact.docx"
        DocumentAssembler().assemble("Exact", [chunk], out)
        doc = DocxDocument(str(out))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert "الفصل" in texts
        assert "المحتوى هنا." in texts

    def test_table_with_extra_paragraphs(self, tmp_path):
        blocks = [
            paragraph("Intro text before the table with some length to it."),
            {"seq": 5, "type": "table", "text": "", "page": 1, "level": 0, "marker": "",
             "rows": [["Term", "Arabic"], ["Low vision", "ضعف البصر"]], "bold": False, "italic": False},
        ]
        # Model returned 3 paragraphs: intro, table rows, trailing note
        translation = (
            "نص المقدمة قبل الجدول.\n\n"
            "المصطلح | الترجمة\nضعف البصر | Low vision\n\n"
            "فقرة إضافية بعد الجدول."
        )
        chunk = FakeChunk(0, "Intro...\n\ntable rows...", translation, blocks)
        out = tmp_path / "tbl.docx"
        DocumentAssembler().assemble("Table Test", [chunk], out)
        doc = DocxDocument(str(out))
        assert len(doc.tables) == 1
        assert doc.tables[0].rows[0].cells[0].text == "المصطلح"
        assert "ضعف البصر" in doc.tables[0].rows[1].cells[0].text
        full = "\n".join(p.text for p in doc.paragraphs)
        assert "نص المقدمة" in full
        assert "فقرة إضافية بعد الجدول." in full
