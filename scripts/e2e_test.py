"""End-to-end workflow test against the RUNNING application (mock provider)."""
import json
import time
import urllib.request

BASE = "http://127.0.0.1:8000/api"
PDF = r"C:\academic-translator\storage\temp\test_book.pdf"


def req(path, method="GET", data=None, is_json=True):
    url = f"{BASE}{path}"
    body = None
    headers = {}
    if data is not None and is_json:
        body = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    r = urllib.request.Request(url, data=body, method=method, headers=headers)
    with urllib.request.urlopen(r, timeout=120) as resp:
        if resp.status == 204:
            return None
        return json.loads(resp.read().decode())


# 1. Upload
boundary = "----adtboundary"
with open(PDF, "rb") as fh:
    file_bytes = fh.read()
parts = []
parts.append(
    f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"Research_Book.pdf\"\r\n"
    f"Content-Type: application/pdf\r\n\r\n".encode()
    + file_bytes + b"\r\n"
)
parts.append(f"--{boundary}--\r\n".encode())
upload_body = b"".join(parts)
r = urllib.request.Request(
    f"{BASE}/documents/upload",
    data=upload_body,
    method="POST",
    headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
)
with urllib.request.urlopen(r, timeout=120) as resp:
    doc = json.loads(resp.read().decode())
print(f"1. Uploaded: {doc['id']} ({doc['original_filename']}, {doc['file_type']}, {doc['file_size']} bytes)")

# 2. Wait for analysis (started automatically after upload? -> analysis runs on job start)
doc = req(f"/documents/{doc['id']}")
print(f"2. Document status: {doc['status']}")

# 3. Start translation
job = req(
    f"/translation/{doc['id']}/start",
    method="POST",
    data={
        "settings": {
            "style": "Academic",
            "domain": "Visual Impairment",
            "use_global_dictionary": True,
            "use_domain_dictionary": True,
            "use_translation_memory": True,
            "chunk_target_chars": 600,
            "use_ocr_if_needed": True,
        }
    },
)
job_id = job["id"]
print(f"3. Job started: {job_id} status={job['status']}")

# 4. Poll progress
final = None
for i in range(240):
    time.sleep(1)
    p = req(f"/translation/jobs/{job_id}/progress")
    if i % 3 == 0 or p["status"] in ("completed", "failed", "paused", "cancelled"):
        print(
            f"   {p['status']:12s} {p['progress_percentage']:5.1f}%  {p['completed_chunks']}/{p['total_chunks']} "
            f"failed={p['failed_chunks']} chapter={p['current_chapter'] or '-'}"
        )
    if p["status"] in ("completed", "failed", "cancelled"):
        final = p
        break
assert final, "job did not finish"
print(f"4. Final status: {final['status']}")

if final["status"] != "completed":
    print("ERROR:", final["error_message"])
    for e in final["recent_errors"]:
        print("  ", e)
    raise SystemExit(1)

# 5. Job detail + quality report
job = req(f"/translation/jobs/{job_id}")
qr = job["quality_report"]
print(
    f"5. Quality: {qr['completed_chunks']}/{qr['total_chunks']} chunks, "
    f"consistent terms={qr['terminology']['consistent_terms']}, "
    f"inconsistent={len(qr['terminology']['inconsistent_terms'])}"
)

# 6. Download DOCX
doc_id = final["document_id"]
url = f"{BASE}/documents/{doc_id}/download/translated"
out_path = r"C:\academic-translator\storage\temp\downloaded_translated.docx"
with urllib.request.urlopen(url, timeout=120) as resp:
    data = resp.read()
with open(out_path, "wb") as fh:
    fh.write(data)
print(f"6. Downloaded DOCX: {out_path} ({len(data)} bytes)")

# 7. Validate DOCX
from docx import Document as DocxDocument
from docx.oxml.ns import qn

d = DocxDocument(out_path)
n_paras = len(d.paragraphs)
rtl_paras = sum(
    1
    for p in d.paragraphs
    if p._p.find(qn("w:pPr")) is not None and p._p.find(qn("w:pPr")).find(qn("w:bidi")) is not None
)
arabic_text = any("\u0600" <= ch <= "\u06FF" for p in d.paragraphs for ch in p.text)
headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
print(f"7. DOCX validation: paragraphs={n_paras}, RTL paragraphs={rtl_paras}, arabic={arabic_text}")
print(f"   headings sample: {headings[:4]}")
print(f"   tables: {len(d.tables)}")
assert rtl_paras > 10, "RTL formatting missing"
assert arabic_text, "No Arabic content"
print("\nE2E WORKFLOW: SUCCESS")
